"""Genera los datos del geoportal a partir de la matriz mensual de UA con amenazas.

Uso (desde la carpeta del repositorio):
    pip install openpyxl pyproj
    python scripts/preparar_ua.py UA_AGO26_AMENAZAS_4.xlsx "Agosto 2026" [UA_AGOSTO_2026]

Genera:
    data/ua.json          datos compactos para el mapa, los puntos y las estadísticas
    data/ua_matriz.json   la matriz completa (todas las columnas, tal cual) para el reporte Excel
    data/diccionario.json pestaña DICCIONARIO_VARIABLES (encabezado, notas y variables)
"""
import sys, json
import openpyxl
from pyproj import Transformer

xlsx, corte = sys.argv[1], sys.argv[2]
hoja = sys.argv[3] if len(sys.argv) > 3 else 'UA_AGOSTO_2026'
wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
ws = wb[hoja]

hdr, raw = None, []
for r in ws.iter_rows(values_only=True):
    if hdr is None:
        if r and r[0] == 'zon_plan':
            hdr = [h for h in r if h is not None]
        continue
    r = list(r[:len(hdr)])
    d = dict(zip(hdr, r))
    if d.get('co_siimdh') is None:
        continue
    raw.append(r)
rows = [dict(zip(hdr, r)) for r in raw]

pars = {f['properties']['DPA_PARROQ'] for f in json.load(open('data/parroquias.json', encoding='utf-8'))['features']}
utm = Transformer.from_crs('EPSG:32717', 'EPSG:4326', always_xy=True)
cat = lambda v: {'1_MEDIA': 1, '2_ALTA': 2, '3_MUY ALTA': 3}.get(v, 0)
cc = lambda r: str(r['cod_cc']).zfill(6)   # parroquia DPA: urbanas agregadas en la cabecera cantonal

d = {k: [] for k in ('serv', 'mod', 'ud', 'tipo', 'area', 'adm')}
def ix(k, v):
    v = v if v not in (None, '') else '-'
    if v not in d[k]: d[k].append(v)
    return d[k].index(v)

out = []
for i, r in enumerate(rows):
    lon, lat = utm.transform(float(r['x']), float(r['y'])) if r['x'] and r['y'] else (None, None)
    out.append([
        str(r['co_siimdh']), r['nombre'], ix('serv', r['servicio']), ix('mod', r['mod']),
        cc(r)[:2], cc(r)[:4], cc(r), r['cabe_canto'],
        ix('ud', r['uni_des']), ix('tipo', r['tipo']), ix('area', r['area']), ix('adm', r['tipo_admin']),
        int(r['total'] or 0),
        cat(r['susc_inundacion']), cat(r['susc_movimientos_masa']), cat(r['susc_sequia']), cat(r['susc_inc_forestales']),
        round(lat, 6) if lat else None, round(lon, 6) if lon else None,
        r['dpa_despro'], r['dpa_descan'], r['dpa_despar'],
    ])

json.dump({'corte': corte, 'dict': d, 'servMod': sorted({(o[2], o[3]) for o in out}), 'rows': out},
          open('data/ua.json', 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))

def clean(v):
    return v if not isinstance(v, float) or not v.is_integer() else int(v)
json.dump({'cols': hdr, 'rows': [[clean(v) for v in r] for r in raw]},
          open('data/ua_matriz.json', 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))

# diccionario
dic = wb['DICCIONARIO_VARIABLES']
lines, variables = [], []
for r in dic.iter_rows(values_only=True):
    vals = [v for v in r if v is not None]
    if not vals: continue
    if isinstance(vals[0], (int, float)) and len(vals) >= 3:
        variables.append([int(vals[0])] + [str(v) for v in vals[1:4]])
    elif vals[0] == 'No. campo' or vals == ['INICIO']:
        continue
    else:
        lines.append([str(v) for v in vals if v != 'INICIO'])
json.dump({'encabezado': lines, 'variables': variables}, open('data/diccionario.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)

miss = sorted({o[6] for o in out if o[6] not in pars})
print(len(out), 'UA |', len(hdr), 'columnas | códigos sin polígono:', miss or 'ninguno', '| sin coordenadas:', sum(1 for o in out if o[17] is None))
